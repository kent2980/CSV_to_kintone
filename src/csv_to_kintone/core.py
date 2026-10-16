"""CSVインポートの業務ロジック。"""
import csv
import json
import os
import tempfile
from datetime import date
from openpyxl import load_workbook
from dataclasses import dataclass, field
from pathlib import Path
from datetime import datetime

@dataclass
class ImportJob:
    id: str
    name: str
    csv_path: str
    app_id: int
    enabled: bool = True
    interval_minutes: int = 60
    field_mapping: dict = field(default_factory=dict)
    duplicate_columns: list = field(default_factory=list)

@dataclass
class ImportResult:
    job_id: str
    status: str = "failed"
    created: int = 0
    updated: int = 0
    deleted: int = 0
    backup_path: str | None = None
    error_message: str | None = None

def read_csv(path: str):
    """CSVヘッダーとデータを読み込み、空データや不正行を拒否します。"""
    with open(path, encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        columns = reader.fieldnames or []
        if not columns or len(set(columns)) != len(columns) or any(not c for c in columns):
            raise ValueError("CSVヘッダーが不正です")
        rows = list(reader)
    if not rows or any(None in row or None in row.values() for row in rows):
        raise ValueError("CSVが空、または不正な行があります")
    return columns, rows

def _cell_text(value) -> str:
    """Excelセルの値をkintone登録用の文字列へ変換します。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_input_file(path: str) -> tuple[list[str], list[dict[str, str]]]:
    """CSVまたはXLSXを読み、列名と文字列化したデータ行を返します。

    Args:
        path: CSVまたはXLSXのパス。XLSXは先頭シート・1行目をヘッダーとして読み込みます。
    Returns:
        (列名一覧, データ行の辞書一覧)。
    Raises:
        ValueError: 未対応形式、空データ、不正ヘッダー、計算結果が未保存の数式。
        OSError: ファイル読み込み失敗時。
    Side effects:
        ファイルを読み込みます。元ファイルは変更しません。
    """
    extension = Path(path).suffix.lower()
    if extension == ".csv":
        return read_csv(path)
    if extension != ".xlsx":
        raise ValueError("未対応のファイル形式です。CSVまたはXLSXを指定してください")
    book = load_workbook(path, read_only=True, data_only=False)
    cached_book = None
    try:
        cached_book = load_workbook(path, read_only=True, data_only=True)
        sheet = book.worksheets[0]
        iterator = sheet.iter_rows()
        cached_iterator = cached_book.worksheets[0].iter_rows(values_only=True)
        header_cells = next(iterator, None)
        header = [cell.value for cell in header_cells] if header_cells else None
        next(cached_iterator, None)
        if not header:
            raise ValueError("Excelヘッダーが不正です")
        columns = [str(value).strip() if value is not None else "" for value in header]
        if not all(columns) or len(columns) != len(set(columns)):
            raise ValueError("Excelヘッダーが空欄または重複しています")
        rows = []
        for number, cells in enumerate(iterator, start=2):
            cached_values = next(cached_iterator, ())
            values = []
            for index, cell in enumerate(cells):
                cached = cached_values[index] if index < len(cached_values) else None
                if cell.data_type == "f":
                    if cached is None:
                        raise ValueError(f"Excelの数式の計算結果が未保存です（{cell.coordinate}）")
                    values.append(_cell_text(cached))
                else:
                    values.append(_cell_text(cell.value))
            if not any(value != "" for value in values):
                continue
            if len(values) > len(columns) and any(value != "" for value in values[len(columns):]):
                raise ValueError(f"Excelの列数が不正です（{number}行目）")
            rows.append(dict(zip(columns, (values + [""] * len(columns))[:len(columns)])))
        if not rows:
            raise ValueError("Excelのデータが0件です。既存レコードは削除しません")
        return columns, rows
    finally:
        book.close()
        if cached_book is not None:
            cached_book.close()


def resolve_field_codes(fields: dict, columns: list, mapping: dict):
    """CSV列名とフィールドラベルを対応付けてコードを返します。"""
    if set(mapping) - set(columns):
        raise ValueError("マッピング対象のCSV列が存在しません")
    result = {}
    supported = {"SINGLE_LINE_TEXT", "MULTI_LINE_TEXT", "NUMBER", "DATE", "DATETIME", "DROP_DOWN", "RADIO_BUTTON", "LINK"}
    for col in columns:
        label = mapping.get(col, col)
        matches = [code for code, prop in fields.items() if prop.get("label") == label]
        if len(matches) != 1 or fields[matches[0]].get("type") not in supported:
            raise ValueError(f"フィールド未発見・重複・未対応: {label}")
        result[col] = matches[0]
    if len(set(result.values())) != len(result):
        raise ValueError("フィールドの割当が重複しています")
    return result

def validate_replace_safety(fields: dict, records: list, codes: dict) -> None:
    """削除前に未対応データ構造と未指定の必須項目を検証します。"""
    unsupported = {"FILE", "SUBTABLE"}
    if any(info.get("type") in unsupported for info in fields.values()):
        raise ValueError("添付ファイルまたはサブテーブルを含むアプリの全件入替は未対応です")
    for record in records:
        if any(value.get("type") in unsupported for value in record.values() if isinstance(value, dict)):
            raise ValueError("添付ファイルまたはサブテーブルを含むレコードの全件入替は未対応です")
    missing = [
        info.get("label", code)
        for code, info in fields.items()
        if info.get("required") and code not in codes.values()
        and info.get("defaultValue") in (None, "")
        and info.get("type") not in {"RECORD_NUMBER", "CREATED_TIME", "UPDATED_TIME", "CREATOR", "MODIFIER"}
    ]
    if missing:
        raise ValueError("CSVに存在しない必須フィールド: " + ", ".join(missing))


def create_backup(app_id: int, records: list, directory: str):
    """削除前の全レコードをJSONに保存し、再読込して検証します。"""
    folder = Path(directory)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"app_{app_id}_{datetime.now():%Y%m%d_%H%M%S_%f}.json"
    content = {"app_id": app_id, "records": records}
    fd, temp_name = tempfile.mkstemp(prefix=".backup_", suffix=".tmp", dir=folder)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(content, stream, ensure_ascii=False)
            stream.flush()
            os.fsync(stream.fileno())
        with open(temp_name, encoding="utf-8") as stream:
            if json.load(stream) != content:
                raise ValueError("バックアップ検証に失敗")
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    return str(path)

def execute_import(job: ImportJob, client, backup_dir: str) -> ImportResult:
    """CSVを検証してUPSERTまたはバックアップ付き全件入替を実行します。"""
    result = ImportResult(job_id=job.id)
    try:
        columns, rows = read_input_file(job.csv_path)
        fields = client.fields(job.app_id)
        codes = resolve_field_codes(fields, columns, job.field_mapping)
        keys = job.duplicate_columns
        if len(keys) != len(set(keys)) or any(key not in columns for key in keys):
            raise ValueError("重複チェック列が不正です")
        payload = [{codes[col]: {"value": row[col]} for col in columns} for row in rows]
        existing = client.records(job.app_id)
        if keys:
            index = {}
            for record in existing:
                key = tuple(str(record[codes[col]]["value"]) for col in keys)
                if key in index:
                    raise ValueError("kintone側に重複キーがあります")
                index[key] = record
            seen, added, updated = set(), [], []
            for row, record in zip(rows, payload):
                key = tuple(row[col] for col in keys)
                if key in seen or any(not value.strip() for value in key):
                    raise ValueError("CSV内に空欄・重複キーがあります")
                seen.add(key)
                if key in index:
                    updated.append({"id": index[key]["$id"]["value"], "record": record})
                else:
                    added.append(record)
            for offset in range(0, len(updated), 100):
                batch = updated[offset:offset + 100]
                client.update(job.app_id, batch)
                result.updated += len(batch)
            for offset in range(0, len(added), 100):
                batch = added[offset:offset + 100]
                client.add(job.app_id, batch)
                result.created += len(batch)
        else:
            validate_replace_safety(fields, existing, codes)
            result.backup_path = create_backup(job.app_id, existing, backup_dir)
            ids = [record["$id"]["value"] for record in existing]
            for offset in range(0, len(ids), 100):
                batch = ids[offset:offset + 100]
                client.delete(job.app_id, batch)
                result.deleted += len(batch)
            for offset in range(0, len(payload), 100):
                batch = payload[offset:offset + 100]
                client.add(job.app_id, batch)
                result.created += len(batch)
        result.status = "success"
    except Exception as error:
        result.error_message = str(error)
    return result
